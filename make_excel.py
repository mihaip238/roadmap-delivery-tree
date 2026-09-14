"""Write a clickable Excel map of Excel KPI items → Jira delivery tree."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\MihaiPostolache\Downloads\kpisss")
SRC = ROOT / "jira_map" / "delivery_tree.json"
OUT = ROOT / "Roadmap_Jira_Delivery_Map.xlsx"
OUT_COPY = ROOT / "jira_map" / "Roadmap_Jira_Delivery_Map.xlsx"

JIRA = "https://eneve.atlassian.net/browse/"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
KIND_FILL = {
    "polaris": PatternFill("solid", fgColor="C6EFCE"),
    "mixed": PatternFill("solid", fgColor="BDD7EE"),
    "inferred": PatternFill("solid", fgColor="FCE4D6"),
    "none": PatternFill("solid", fgColor="F2F2F2"),
}


def browse(key: str | None) -> str:
    if not key:
        return ""
    return JIRA + key


def write_header(ws, headers: list[str]) -> None:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def set_widths(ws, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def link_cell(cell, key: str | None) -> None:
    if not key:
        cell.value = ""
        return
    cell.value = key
    cell.hyperlink = browse(key)
    cell.font = LINK_FONT


def style_row(ws, row: int, cols: int, kind: str | None = None) -> None:
    fill = KIND_FILL.get(kind or "")
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.alignment = WRAP
        cell.border = THIN
        if fill and col <= cols:
            if cell.font and cell.font.color and cell.font.underline:
                continue
            if not cell.hyperlink:
                cell.fill = fill


def main() -> None:
    data = json.loads(SRC.read_text(encoding="utf-8"))
    all_items = list(data.get("items") or []) + list(data.get("non_ebase") or [])

    wb = Workbook()

    # --- Read me ---
    intro = wb.active
    intro.title = "Read_me"
    intro["A1"] = "Roadmap Jira delivery map"
    intro["A1"].font = Font(bold=True, size=16)
    lines = [
        "",
        "This workbook is the stored mapping of every KPI Excel item to Jira.",
        "Polaris = ground-truth EDP→EPP link. Inferred = title match, review before using in KPIs. None = no honest connection.",
        "",
        "Sheets:",
        "  Items     — one row per Excel item (94 EBASE + 31 Non-EBASE). Click the EDP key to open Jira.",
        "  EDP_EPP   — one row per Excel item × EPP (or JTP epic). Both keys are clickable.",
        "  Tree      — full walk: EDP → EPP → Feature → Story. Every key is a Jira hyperlink. URL columns are there to copy.",
        "  Active_32 — Now / Next / Later EBASE items only.",
        "  Time_EDP  — rolled vs unique hours per EDP. Unique de-duplicates shared EPPs.",
        "",
        "Time is native Jira worklogs. 1d = 8h. Feature rolled time uses Jira aggregatetimespent",
        "so defects under a feature still count even if they were missing from an older dump.",
        "EPP rolled time = worklogs on the epic + rolled time of its child features.",
        "An EPP linked to two EDPs is shown on both; unique hours count each ticket once.",
        "",
        "Jira site: https://eneve.atlassian.net",
        "Issue URL pattern: https://eneve.atlassian.net/browse/<KEY>",
        "",
        "Source JSON: jira_map/delivery_tree.json",
        "Walk date: 11 Sep 2026. Story counts include Tasks, Questions, and Xray tests under features.",
    ]
    for i, line in enumerate(lines, 2):
        intro.cell(i, 1, line)
        intro.cell(i, 1).alignment = WRAP
    intro.column_dimensions["A"].width = 140
    intro.row_dimensions[3].height = 36

    # --- Items ---
    items_ws = wb.create_sheet("Items")
    item_headers = [
        "family", "excel_summary", "product", "roadmap", "active",
        "customer_problem", "origin", "business_case", "planned_pi",
        "link_kind", "edp_key", "edp_url", "edp_status",
        "epp_count", "feature_count", "story_count",
        "spent_hours", "unique_spent_hours", "estimate_hours", "spent_pretty", "estimate_pretty",
        "epp_keys", "epp_urls",
    ]
    write_header(items_ws, item_headers)
    for r, item in enumerate(all_items, 2):
        excel = item.get("excel") or {}
        edp = item.get("edp") or {}
        cov = item.get("coverage") or {}
        epps = item.get("epps") or []
        keys = [e.get("key") for e in epps if e.get("key")]
        kind = item.get("link_kind") or ""
        t = (edp.get("time") or {})
        values = [
            item.get("family") or excel.get("family") or "",
            excel.get("summary") or "",
            excel.get("product") or "",
            excel.get("roadmap") or "",
            "TRUE" if excel.get("active") else "FALSE",
            excel.get("customer_problem") or "",
            excel.get("origin") or "",
            excel.get("business_case") or "",
            excel.get("planned_pi") or "",
            kind,
            edp.get("key") or "",
            browse(edp.get("key")),
            edp.get("status") or "",
            cov.get("epp_count") or len(epps),
            cov.get("feature_count") or 0,
            cov.get("story_count") or 0,
            t.get("rolledSpentHours") if t.get("rolledSpentHours") is not None else cov.get("spentHours") or 0,
            t.get("uniqueSpentHours") if t.get("uniqueSpentHours") is not None else cov.get("uniqueSpentHours") or 0,
            t.get("rolledEstimateHours") if t.get("rolledEstimateHours") is not None else "",
            t.get("rolledSpent") or "",
            t.get("rolledEstimate") or "",
            ", ".join(keys),
            "\n".join(browse(k) for k in keys),
        ]
        for c, val in enumerate(values, 1):
            items_ws.cell(r, c, val)
        link_cell(items_ws.cell(r, 11), edp.get("key"))
        items_ws.cell(r, 12).hyperlink = browse(edp.get("key")) if edp.get("key") else None
        if edp.get("key"):
            items_ws.cell(r, 12).font = LINK_FONT
        fill = KIND_FILL.get(kind)
        if fill:
            items_ws.cell(r, 10).fill = fill
        for c in range(1, len(item_headers) + 1):
            items_ws.cell(r, c).alignment = WRAP
            items_ws.cell(r, c).border = THIN
    items_ws.auto_filter.ref = f"A1:{get_column_letter(len(item_headers))}{1 + len(all_items)}"
    set_widths(items_ws, {
        1: 12, 2: 48, 3: 18, 4: 10, 5: 10, 6: 16, 7: 16, 8: 14, 9: 18,
        10: 12, 11: 14, 12: 42, 13: 14, 14: 12, 15: 14, 16: 12,
        17: 14, 18: 16, 19: 14, 20: 16, 21: 16, 22: 36, 23: 50,
    })

    # --- EDP_EPP ---
    pair_ws = wb.create_sheet("EDP_EPP")
    pair_headers = [
        "family", "excel_summary", "roadmap", "active", "link_kind",
        "edp_key", "edp_url", "edp_status",
        "epp_key", "epp_url", "epp_summary", "epp_status", "epp_method", "epp_score",
        "feature_count", "story_count",
        "spent_hours", "estimate_hours", "spent_pretty", "shared_with",
    ]
    write_header(pair_ws, pair_headers)
    pr = 2
    for item in all_items:
        excel = item.get("excel") or {}
        edp = item.get("edp") or {}
        kind = item.get("link_kind") or ""
        epps = item.get("epps") or [None]
        if not epps:
            epps = [None]
        for epp in epps:
            epp = epp or {}
            epp_time = epp.get("time") or {}
            row = [
                item.get("family") or excel.get("family") or "",
                excel.get("summary") or "",
                excel.get("roadmap") or "",
                "TRUE" if excel.get("active") else "FALSE",
                kind,
                edp.get("key") or "",
                browse(edp.get("key")),
                edp.get("status") or "",
                epp.get("key") or "",
                browse(epp.get("key")),
                epp.get("summary") or "",
                epp.get("status") or "",
                epp.get("method") or "",
                epp.get("score") if epp.get("score") is not None else "",
                epp.get("feature_count") if epp else 0,
                epp.get("story_count") if epp else 0,
                epp_time.get("rolledSpentHours") if epp_time.get("rolledSpentHours") is not None else 0,
                epp_time.get("rolledEstimateHours") if epp_time.get("rolledEstimateHours") is not None else "",
                epp_time.get("rolledSpent") or "",
                ", ".join(epp_time.get("sharedWith") or []),
            ]
            for c, val in enumerate(row, 1):
                pair_ws.cell(pr, c, val)
                pair_ws.cell(pr, c).alignment = WRAP
                pair_ws.cell(pr, c).border = THIN
            link_cell(pair_ws.cell(pr, 6), edp.get("key"))
            if edp.get("key"):
                pair_ws.cell(pr, 7).hyperlink = browse(edp.get("key"))
                pair_ws.cell(pr, 7).font = LINK_FONT
            link_cell(pair_ws.cell(pr, 9), epp.get("key"))
            if epp.get("key"):
                pair_ws.cell(pr, 10).hyperlink = browse(epp.get("key"))
                pair_ws.cell(pr, 10).font = LINK_FONT
            fill = KIND_FILL.get(kind)
            if fill:
                pair_ws.cell(pr, 5).fill = fill
            pr += 1
    pair_ws.auto_filter.ref = f"A1:{get_column_letter(len(pair_headers))}{pr - 1}"
    set_widths(pair_ws, {
        1: 12, 2: 44, 3: 10, 4: 10, 5: 12, 6: 14, 7: 42, 8: 14,
        9: 14, 10: 42, 11: 48, 12: 16, 13: 16, 14: 12, 15: 14, 16: 12,
        17: 14, 18: 14, 19: 16, 20: 28,
    })

    # --- Tree ---
    tree_ws = wb.create_sheet("Tree")
    tree_headers = [
        "family", "excel_summary", "product", "roadmap", "active", "link_kind",
        "edp_key", "edp_url", "edp_status",
        "epp_key", "epp_url", "epp_summary", "epp_status", "epp_method", "epp_score",
        "feature_key", "feature_url", "feature_summary", "feature_type", "feature_status",
        "feature_spent_hours", "feature_estimate_hours",
        "story_key", "story_url", "story_summary", "story_type", "story_status",
        "story_spent_hours", "story_estimate_hours", "story_points",
    ]
    write_header(tree_ws, tree_headers)
    tr = 2
    for item in all_items:
        excel = item.get("excel") or {}
        edp = item.get("edp") or {}
        kind = item.get("link_kind") or ""
        epps = item.get("epps") or []
        if not epps:
            epps = [{}]
        for epp in epps:
            children = epp.get("children") or []
            if not children:
                children = [{}]
            for feat in children:
                stories = feat.get("stories") or []
                if not stories:
                    stories = [{}]
                for st in stories:
                    feat_time = feat.get("time") or {}
                    st_time = st.get("time") or {}
                    row = [
                        item.get("family") or excel.get("family") or "",
                        excel.get("summary") or "",
                        excel.get("product") or "",
                        excel.get("roadmap") or "",
                        "TRUE" if excel.get("active") else "FALSE",
                        kind,
                        edp.get("key") or "",
                        browse(edp.get("key")),
                        edp.get("status") or "",
                        epp.get("key") or "",
                        browse(epp.get("key")),
                        epp.get("summary") or "",
                        epp.get("status") or "",
                        epp.get("method") or "",
                        epp.get("score") if epp.get("score") is not None else "",
                        feat.get("key") or "",
                        browse(feat.get("key")),
                        feat.get("summary") or "",
                        feat.get("issuetype") or "",
                        feat.get("status") or "",
                        feat_time.get("rolledSpentHours") if feat_time.get("rolledSpentHours") is not None else "",
                        feat_time.get("rolledEstimateHours") if feat_time.get("rolledEstimateHours") is not None else "",
                        st.get("key") or "",
                        browse(st.get("key")),
                        st.get("summary") or "",
                        st.get("issuetype") or "",
                        st.get("status") or "",
                        st_time.get("rolledSpentHours") if st_time.get("rolledSpentHours") is not None else "",
                        st_time.get("rolledEstimateHours") if st_time.get("rolledEstimateHours") is not None else "",
                        st_time.get("rolledPoints") if st_time.get("rolledPoints") is not None else "",
                    ]
                    for c, val in enumerate(row, 1):
                        tree_ws.cell(tr, c, val)
                        tree_ws.cell(tr, c).alignment = WRAP
                        tree_ws.cell(tr, c).border = THIN
                    for col, key in ((7, edp.get("key")), (10, epp.get("key")), (16, feat.get("key")), (23, st.get("key"))):
                        link_cell(tree_ws.cell(tr, col), key)
                    for col, key in ((8, edp.get("key")), (11, epp.get("key")), (17, feat.get("key")), (24, st.get("key"))):
                        if key:
                            tree_ws.cell(tr, col).hyperlink = browse(key)
                            tree_ws.cell(tr, col).font = LINK_FONT
                    fill = KIND_FILL.get(kind)
                    if fill:
                        tree_ws.cell(tr, 6).fill = fill
                    tr += 1
    tree_ws.auto_filter.ref = f"A1:{get_column_letter(len(tree_headers))}{tr - 1}"
    set_widths(tree_ws, {
        1: 12, 2: 36, 3: 16, 4: 10, 5: 10, 6: 12,
        7: 14, 8: 42, 9: 14,
        10: 14, 11: 42, 12: 40, 13: 14, 14: 16, 15: 10,
        16: 16, 17: 42, 18: 44, 19: 16, 20: 16, 21: 14, 22: 14,
        23: 16, 24: 42, 25: 44, 26: 14, 27: 14, 28: 14, 29: 14, 30: 12,
    })
    tree_ws.sheet_properties.tabColor = "1F4E79"

    time_ws = wb.create_sheet("Time_EDP")
    time_headers = [
        "edp_key", "edp_url", "excel_summary", "product", "roadmap", "active", "link_kind",
        "spent_hours", "unique_spent_hours", "estimate_hours", "spent_pretty", "estimate_pretty",
        "story_points", "shared_epps", "feature_count", "story_count",
    ]
    write_header(time_ws, time_headers)
    tr_edp = 2
    for item in all_items:
        excel = item.get("excel") or {}
        edp = item.get("edp") or {}
        t = edp.get("time") or {}
        cov = item.get("coverage") or {}
        if not edp.get("key") and not (t.get("rolledSpentHours") or 0):
            continue
        row = [
            edp.get("key") or "",
            browse(edp.get("key")),
            excel.get("summary") or "",
            excel.get("product") or "",
            excel.get("roadmap") or "",
            "TRUE" if excel.get("active") else "FALSE",
            item.get("link_kind") or "",
            t.get("rolledSpentHours") or 0,
            t.get("uniqueSpentHours") or 0,
            t.get("rolledEstimateHours") if t.get("rolledEstimateHours") is not None else "",
            t.get("rolledSpent") or "",
            t.get("rolledEstimate") or "",
            t.get("rolledPoints") if t.get("rolledPoints") is not None else "",
            ", ".join(t.get("sharedWith") or []),
            cov.get("feature_count") or 0,
            cov.get("story_count") or 0,
        ]
        for c, val in enumerate(row, 1):
            time_ws.cell(tr_edp, c, val)
            time_ws.cell(tr_edp, c).alignment = WRAP
            time_ws.cell(tr_edp, c).border = THIN
        link_cell(time_ws.cell(tr_edp, 1), edp.get("key"))
        if edp.get("key"):
            time_ws.cell(tr_edp, 2).hyperlink = browse(edp.get("key"))
            time_ws.cell(tr_edp, 2).font = LINK_FONT
        tr_edp += 1
    time_ws.auto_filter.ref = f"A1:{get_column_letter(len(time_headers))}{max(tr_edp - 1, 2)}"
    set_widths(time_ws, {
        1: 14, 2: 42, 3: 48, 4: 18, 5: 10, 6: 10, 7: 12,
        8: 14, 9: 16, 10: 14, 11: 16, 12: 16, 13: 12, 14: 28, 15: 14, 16: 12,
    })
    time_ws.sheet_properties.tabColor = "833C0C"

    # --- Active_32 ---
    act_ws = wb.create_sheet("Active_32")
    write_header(act_ws, item_headers)
    ar = 2
    for item in data.get("items") or []:
        excel = item.get("excel") or {}
        if not excel.get("active"):
            continue
        edp = item.get("edp") or {}
        cov = item.get("coverage") or {}
        epps = item.get("epps") or []
        keys = [e.get("key") for e in epps if e.get("key")]
        kind = item.get("link_kind") or ""
        t = edp.get("time") or {}
        values = [
            item.get("family") or "",
            excel.get("summary") or "",
            excel.get("product") or "",
            excel.get("roadmap") or "",
            "TRUE",
            excel.get("customer_problem") or "",
            excel.get("origin") or "",
            excel.get("business_case") or "",
            excel.get("planned_pi") or "",
            kind,
            edp.get("key") or "",
            browse(edp.get("key")),
            edp.get("status") or "",
            cov.get("epp_count") or len(epps),
            cov.get("feature_count") or 0,
            cov.get("story_count") or 0,
            t.get("rolledSpentHours") if t.get("rolledSpentHours") is not None else 0,
            t.get("uniqueSpentHours") if t.get("uniqueSpentHours") is not None else 0,
            t.get("rolledEstimateHours") if t.get("rolledEstimateHours") is not None else "",
            t.get("rolledSpent") or "",
            t.get("rolledEstimate") or "",
            ", ".join(keys),
            "\n".join(browse(k) for k in keys),
        ]
        for c, val in enumerate(values, 1):
            act_ws.cell(ar, c, val)
            act_ws.cell(ar, c).alignment = WRAP
            act_ws.cell(ar, c).border = THIN
        link_cell(act_ws.cell(ar, 11), edp.get("key"))
        if edp.get("key"):
            act_ws.cell(ar, 12).hyperlink = browse(edp.get("key"))
            act_ws.cell(ar, 12).font = LINK_FONT
        fill = KIND_FILL.get(kind)
        if fill:
            act_ws.cell(ar, 10).fill = fill
        ar += 1
    act_ws.auto_filter.ref = f"A1:{get_column_letter(len(item_headers))}{ar - 1}"
    set_widths(act_ws, {
        1: 12, 2: 48, 3: 18, 4: 10, 5: 10, 6: 16, 7: 16, 8: 14, 9: 18,
        10: 12, 11: 14, 12: 42, 13: 14, 14: 12, 15: 14, 16: 12,
        17: 14, 18: 16, 19: 14, 20: 16, 21: 16, 22: 36, 23: 50,
    })

    # --- Gaps ---
    gap_ws = wb.create_sheet("Gaps")
    gap_headers = [
        "family", "excel_summary", "roadmap", "active", "link_kind",
        "edp_key", "edp_url", "note",
    ]
    write_header(gap_ws, gap_headers)
    gr = 2
    for item in all_items:
        excel = item.get("excel") or {}
        kind = item.get("link_kind") or ""
        cov = item.get("coverage") or {}
        is_gap = kind == "none" or (
            excel.get("active") and (cov.get("feature_count") or 0) == 0 and (cov.get("story_count") or 0) == 0
        )
        if not is_gap:
            continue
        edp = item.get("edp") or {}
        note = item.get("note") or ""
        if kind != "none" and not note:
            note = "Mapped to an EPP, but that EPP currently has no Feature children in Jira."
        row = [
            item.get("family") or excel.get("family") or "",
            excel.get("summary") or "",
            excel.get("roadmap") or "",
            "TRUE" if excel.get("active") else "FALSE",
            kind,
            edp.get("key") or "",
            browse(edp.get("key")),
            note,
        ]
        for c, val in enumerate(row, 1):
            gap_ws.cell(gr, c, val)
            gap_ws.cell(gr, c).alignment = WRAP
            gap_ws.cell(gr, c).border = THIN
        link_cell(gap_ws.cell(gr, 6), edp.get("key"))
        if edp.get("key"):
            gap_ws.cell(gr, 7).hyperlink = browse(edp.get("key"))
            gap_ws.cell(gr, 7).font = LINK_FONT
        fill = KIND_FILL.get(kind)
        if fill:
            gap_ws.cell(gr, 5).fill = fill
        gr += 1
    gap_ws.auto_filter.ref = f"A1:{get_column_letter(len(gap_headers))}{max(gr - 1, 2)}"
    set_widths(gap_ws, {1: 12, 2: 52, 3: 10, 4: 10, 5: 12, 6: 14, 7: 42, 8: 70})

    intro.sheet_properties.tabColor = "833C0C"
    items_ws.sheet_properties.tabColor = "1F4E79"
    pair_ws.sheet_properties.tabColor = "2E75B6"
    act_ws.sheet_properties.tabColor = "548235"
    gap_ws.sheet_properties.tabColor = "C00000"

    wb.save(OUT)
    wb.save(OUT_COPY)
    print(f"wrote {OUT}")
    print(f"copy {OUT_COPY}")
    print(f"items {len(all_items)} tree_rows {tr - 2} pairs {pr - 2} active {ar - 2} gaps {gr - 2}")


if __name__ == "__main__":
    main()
